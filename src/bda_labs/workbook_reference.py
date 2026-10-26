"""Read cached values and expanded formulas once; never evaluate untrusted workbook formulas."""
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET
import json,zipfile
import openpyxl
from openpyxl.utils.cell import range_boundaries,get_column_letter


@dataclass
class Sheet:
    name:str
    values:dict
    formulas:dict
    errors:dict

    def get(self,address,default=None):return self.values.get(address,default)

    def matrix(self,address):
        c1,r1,c2,r2=range_boundaries(address)
        return [[self.get(f'{get_column_letter(c)}{r}') for c in range(c1,c2+1)] for r in range(r1,r2+1)]


def read_cached_book(path):
    path=Path(path);sheets={}
    # A streaming pass per mode avoids quadratic repeated ReadOnlyWorksheet[cell] scans.
    for data_only in (True,False):
        wb=openpyxl.load_workbook(path,read_only=True,data_only=data_only)
        try:
            for ws in wb:
                if data_only:sheets[ws.title]=Sheet(ws.title,{}, {}, {})
                target=sheets[ws.title]
                ws.reset_dimensions()
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value is None:continue
                        if data_only:
                            target.values[cell.coordinate]=cell.value
                            if cell.data_type=='e':target.errors[cell.coordinate]=cell.value
                        elif cell.data_type=='f':target.formulas[cell.coordinate]=str(cell.value)
        finally:wb.close()
    return sheets


def coverage_table(books,descriptions,path):
    rows=['# 工作表迁移对照','','原工作簿只读；缓存是教材保存的快照，不等于本次重算。','','| 来源键 | 工作表 | 教学内容与Python去向 |','|---|---|---|']
    expected={(key,sheet) for key,book in books.items() for sheet in book}
    if set(descriptions)!=expected:raise AssertionError(f'Workbook coverage differs: {expected^set(descriptions)}')
    for key,sheet in sorted(expected):rows.append(f'| {key} | {sheet} | {descriptions[(key,sheet)]} |')
    path.write_text('\n'.join(rows)+'\n',encoding='utf-8')
    return {'source_sheets':len(expected),'mapped_sheets':len(descriptions)}


def pbix_visuals(path):
    with zipfile.ZipFile(path) as z:layout=json.loads(z.read('Report/Layout').decode('utf-16le'))
    result=[]
    for page in layout['sections']:
        for number,container in enumerate(page.get('visualContainers',[]),1):
            config=json.loads(container.get('config','{}'));visual=config.get('singleVisual',{})
            query=visual.get('prototypeQuery',{})
            result.append({'page':page.get('displayName',page.get('name')),'visual':number,'kind':visual.get('visualType','unknown'),
                           'fields':'; '.join(v.get('queryRef','') for values in visual.get('projections',{}).values() for v in values),
                           'aggregations':[{'reference_name':s.get('Name'),'display_name':s.get('NativeReferenceName'),
                                            'function_code':s['Aggregation']['Function']}
                                           for s in query.get('Select',[]) if 'Aggregation' in s]})
    return result
