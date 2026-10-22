"""B08: three textbook funnels with distinct step and cumulative denominators."""
from __future__ import annotations
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from .common import context,finish,plotting,save_table,compare_numeric


def calculate(stages,counts):
    stages=list(stages)
    values=np.asarray(counts,dtype=float)
    if len(stages)!=len(values) or len(values)==0:raise ValueError('Stages and counts must be nonempty and aligned')
    if len(set(stages))!=len(stages) or any(not str(s).strip() for s in stages):raise ValueError('Stage labels must be unique and nonempty')
    if not np.isfinite(values).all() or (values<0).any() or (values%1!=0).any():raise ValueError('User counts must be finite nonnegative integers')
    if (np.diff(values)>0).any():raise ValueError('Counts increase; review cohort/stage comparability before calling this a conversion funnel')
    previous=np.r_[values[0],values[:-1]]
    step=np.divide(values,previous,out=np.full_like(values,np.nan),where=previous!=0)
    cumulative=values/values[0] if values[0]>0 else np.full_like(values,np.nan)
    return pd.DataFrame({'stage':stages,'users':values.astype('int64'),'step_conversion':step,
        'cumulative_conversion':cumulative,'lost_from_previous':(previous-values).astype('int64'),
        'chart_padding':(values[0]-values)/2})


def main(argv=None):
    ctx=context('B08',argv)
    wb=load_workbook(ctx.source('funnel_raw'),read_only=True,data_only=True)
    demo=list(wb['演示'].iter_rows(min_row=4,max_row=9,max_col=3,values_only=True))
    exercise=list(wb['练习'].iter_rows(max_row=15,max_col=4,values_only=True))
    cases={
        '销售漏斗':calculate([r[0] for r in demo],[r[1] for r in demo]),
        '购买行为漏斗':calculate([r[1] for r in exercise[2:6]],[r[2] for r in exercise[2:6]]),
        'AARRR教学示例':calculate([r[1] for r in exercise[10:15]],[r[2] for r in exercise[10:15]])}
    wb.close()
    wb=load_workbook(ctx.source('funnel_reference'),read_only=True,data_only=True)
    ref=list(wb['Sheet1'].iter_rows(min_row=4,max_row=9,max_col=5,values_only=True))
    wb.close()
    if cases['销售漏斗'].stage.tolist()!=[r[0] for r in ref]:raise AssertionError('Reference funnel stages differ')
    expected=np.array([[r[2],r[3],r[4],r[1]] for r in ref],dtype=float)
    checks=compare_numeric(cases['销售漏斗'][['users','step_conversion','cumulative_conversion','chart_padding']],expected,'Sales funnel reference')
    combined=pd.concat([frame.assign(case=name) for name,frame in cases.items()],ignore_index=True)
    save_table(combined,ctx.output/'funnel_metrics.csv')
    plt=plotting()
    fig,axes=plt.subplots(1,3,figsize=(16,6),layout='constrained')
    for ax,(name,frame) in zip(axes,cases.items()):
        ax.barh(frame.stage,frame.users,left=frame.chart_padding,color='#087f8c',alpha=.9,height=.7)
        ax.invert_yaxis()
        ax.set_xlim(0,frame.users.iloc[0]*1.02)
        ax.set_title(name,fontsize=13)
        ax.set_xlabel('示例阶段人数')
        for i,row in frame.iterrows():
            ax.text(frame.users.iloc[0]/2,i,f'{row.users}人',ha='center',va='center',color='white',fontsize=11)
            ax.text(1.03,i,f'逐级 {row.step_conversion:.1%}\n累计 {row.cumulative_conversion:.1%}',
                transform=ax.get_yaxis_transform(),va='center',fontsize=8,color='#466070')
        ax.tick_params(axis='x',labelsize=9)
    fig.suptitle('漏斗分析：逐级转化与累计转化使用不同分母',fontsize=17,color='#17364d')
    fig.savefig(ctx.output/'funnel_analysis.png',dpi=150,bbox_inches='tight')
    plt.close(fig)
    results=[]
    for name,frame in cases.items():
        results.append({'case':name,'stages':len(frame),'first_stage_users':int(frame.users.iloc[0]),
            'last_stage_users':int(frame.users.iloc[-1]),'overall_conversion':float(frame.cumulative_conversion.iloc[-1])})
    summary={'lab_id':'B08','status':'reference_verified','reference_scope':'Sales demo: stage users, step/cumulative conversion and invisible chart padding; practice funnels calculated without an answer sheet.',
        'reference_checks':checks,'cases':results,'zero_denominator_policy':'unavailable, never silently zero',
        'data_scope':'Aggregated teaching counts; no event-level paths or cohort membership supplied.'}
    report="""# B08 转化漏斗

## 三组教学结果

- 销售漏斗：100名潜在用户到50名成交用户，累计转化50%。
- 购买行为：100名示例浏览用户到80名购买用户，累计转化80%。
- AARRR示例：100名获取用户到30名推荐用户，累计比例30%。

![三组漏斗](funnel_analysis.png)

## 计算方法

逐级转化=本阶段人数/上一阶段人数；累计转化=本阶段人数/首阶段人数。比如销售漏斗的“意向”阶段为80/90=88.89%，累计是80/100=80%，两者不能混用。首阶段在分母非零时为100%。分母为0时明确保留不可用，不把0/0填为0%。

居中漏斗使用(首阶段人数-本阶段人数)/2作为绘图占位，对应Excel的透明堆积辅助列。

## 参考核对与限制

销售漏斗的6个阶段，对人数、逐级转化、总体转化和绘图占位共24项逐项核对参考缓存。另两组练习只有原始阶段人数，没有配套答案，不声称它们也通过参考答案核对。

资料是汇总教学人数，没有逐用户事件路径。递减数量不能证明每一级都由上一级同一批人进入，AARRR也不是天然严格嵌套的路径。因此这些比率是给定示例的算术展示，不能用来推断真实客户的逐级条件转化或归因。实际业务需要核对用户集合、阶段顺序、去重和观察窗口。
"""
    finish(ctx,summary,report)


if __name__=='__main__':main()
