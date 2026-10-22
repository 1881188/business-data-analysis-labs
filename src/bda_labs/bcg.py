"""B04: textbook-specific BCG using two-month sales share and month-on-month growth."""
from __future__ import annotations
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from .common import context,finish,plotting,save_table,compare_numeric
from .rfm import money_cents

COMPANY=set('EFGHIJ')


def quadrant(share,growth,share_cutoff,growth_cutoff):
    if not np.isfinite(growth):return '无法比较'
    high_share=share>=share_cutoff
    high_growth=growth>=growth_cutoff
    return {(True,True):'明星产品',(True,False):'金牛产品',
            (False,True):'问题产品',(False,False):'瘦狗产品'}[(high_share,high_growth)]


def calculate(raw,months=('2021-07','2021-08')):
    df=raw.copy()
    if df[['订单ID','销售日期','品牌','单价','销量']].isna().any().any():raise ValueError('Missing sales fields')
    if df['订单ID'].duplicated().any():raise ValueError('Expected one row per order in this case')
    quantity=pd.to_numeric(df['销量'],errors='raise')
    if not np.isfinite(quantity).all() or (quantity<0).any() or (quantity%1!=0).any():raise ValueError('Quantity must be a nonnegative integer')
    if (df['单价']<0).any():raise ValueError('Unexpected negative unit price')
    df['amount_cents']=money_cents(df['单价'])*quantity.astype('int64')
    df['month']=pd.to_datetime(df['销售日期'],errors='raise').dt.to_period('M').astype(str)
    if set(df['month'])!=set(months):raise ValueError('Expected exactly the specified two months')
    table=df.pivot_table(index='品牌',columns='month',values='amount_cents',aggfunc='sum',fill_value=0).reindex(columns=months)
    table.columns=['previous_cents','current_cents']
    previous,current=int(table.previous_cents.sum()),int(table.current_cents.sum())
    if previous<=0 or current<=0:raise ValueError('Both market-period totals must be positive')
    table['combined_cents']=table.previous_cents+table.current_cents
    table['sales_share']=table.combined_cents/(previous+current)
    table['growth']=table.current_cents.div(table.previous_cents.replace(0,np.nan))-1
    growth_cutoff=current/previous-1
    share_cutoff=float(table.sales_share.mean())
    table['quadrant']=[quadrant(s,g,share_cutoff,growth_cutoff) for s,g in zip(table.sales_share,table.growth)]
    table['company_brand']=table.index.isin(COMPANY)
    return table.rename_axis('brand').reset_index(),{'previous_cents':previous,'current_cents':current,
        'market_growth':growth_cutoff,'average_share':share_cutoff}


def main(argv=None):
    ctx=context('B04',argv)
    raw=pd.read_excel(ctx.source('bcg_raw'),sheet_name='销售表',usecols='A:F')
    brands,totals=calculate(raw)
    wb=load_workbook(ctx.source('bcg_reference'),read_only=True,data_only=True)
    rows=list(wb['参考答案'].iter_rows(max_row=21,max_col=12,values_only=True))
    ref=pd.DataFrame([r[9:12] for r in rows[5:19]],columns=['brand','sales_share','growth']).set_index('brand')
    indexed=brands.set_index('brand').reindex(ref.index)
    if set(brands.brand)!=set(ref.index):raise AssertionError('Reference brand sets differ')
    checks={'brand_metrics':compare_numeric(indexed[['sales_share','growth']],ref,'BCG brand metrics'),
        'thresholds':compare_numeric([totals['market_growth'],totals['average_share']],[rows[0][9],rows[1][9]],'BCG thresholds')}
    wb.close()
    save_table(brands,ctx.output/'brands.csv')
    plt=plotting()
    from matplotlib.lines import Line2D
    fig,ax=plt.subplots(figsize=(11,7),layout='constrained')
    ax.axvline(totals['average_share']*100,color='#96a6af',linestyle='--')
    ax.axhline(totals['market_growth']*100,color='#96a6af',linestyle='--')
    for _,r in brands.iterrows():
        color='#087f8c' if r.company_brand else '#9caeba'
        ax.scatter(r.sales_share*100,r.growth*100,s=90 if r.company_brand else 55,color=color,zorder=3)
        ax.annotate(r.brand,(r.sales_share*100,r.growth*100),xytext=(5,5),textcoords='offset points',fontsize=11)
    ax.set(title='波士顿矩阵：两月合计份额与八月环比',xlabel='七至八月合计销售额占比（%）',ylabel='八月较七月销售额增长率（%）')
    ax.legend(handles=[Line2D([],[],marker='o',linestyle='',color='#087f8c',label='本公司E-J品牌'),
                       Line2D([],[],marker='o',linestyle='',color='#9caeba',label='其他品牌')],
              title=f"分界：份额{totals['average_share']:.2%} / 增长{totals['market_growth']:.2%}",
              loc='upper right',frameon=False,fontsize=9,title_fontsize=9)
    fig.savefig(ctx.output/'bcg_matrix.png',dpi=150)
    plt.close(fig)
    company=brands.loc[brands.company_brand,['brand','quadrant']]
    labels='、'.join(f'{r.brand}：{r.quadrant}' for r in company.itertuples())
    summary={'lab_id':'B04','status':'reference_verified','source_rows':len(raw),'brands':len(brands),
        'months':['2021-07','2021-08'],'share_basis':'combined_sales_of_both_months',
        'threshold_equality_policy':'high','market':totals,'reference_checks':checks,
        'company_classification':company.to_dict('records')}
    report=f"""# B04 波士顿矩阵

## 结果与业务问题

原始{len(raw):,}笔销售，覆盖{len(brands)}个品牌。七月市场销售额{totals['previous_cents']/100:,.2f}元，八月{totals['current_cents']/100:,.2f}元，市场环比{totals['market_growth']:.4%}。本公司有E、F、G、H、I、J六个品牌。

![波士顿矩阵](bcg_matrix.png)

按本例阈值：{labels}。这是一组销售位置标签，无法单凭两月销售判断利润或资金回收，不应直接据此退出某个品牌。

## 教材迁移口径

单笔金额=单价×销量，先转整数分再汇总。横轴为七八月合计销售额/两月全市场销售额；纵轴为八月/七月销售额-1。不同规格销量不直接相加比较市场份额。

份额分界为14个品牌的平均份额{totals['average_share']:.4%}，增长分界为市场环比{totals['market_growth']:.4%}。这遵循本教材的销售额份额定义，不改写成“相对最大竞争对手份额”。等于阈值归高是本实现明确约定；教材没有边界样本可供标签核对。

## 核对与限制

14个品牌的份额和环比共28项，以及两条数值阈值，与参考答案缓存一致（浮点容差1e-12）。参考未提供可读取的逐品牌象限标签，标签依据公开说明的规则计算，不能声称该部分也逐标签对账。

若前期销售为0，增长率不可比较，标“无法比较”，不填0或无穷大。当前数据只覆盖两个月，不包含利润、推广成本和长期市场潜力。完整聚合见brands.csv，原始材料保持只读。
"""
    finish(ctx,summary,report)


if __name__=='__main__':main()
