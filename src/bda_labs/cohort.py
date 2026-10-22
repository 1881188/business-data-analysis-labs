"""B07: daily cohort retention with explicit observation maturity and duplicate-login control."""
from __future__ import annotations
from datetime import datetime
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from .common import context,finish,plotting,save_table,compare_numeric

OBSERVATION_END=pd.Timestamp('2020-07-31')


def calculate(raw,observation_end,horizon_days=31):
    if horizon_days<1:raise ValueError('Retention horizon must be positive')
    df=raw[['用户编号','注册日期','登录时点']].copy()
    if df.isna().any().any():raise ValueError('Missing user, registration or login')
    df['注册日期']=pd.to_datetime(df['注册日期'],errors='raise').dt.normalize()
    df['login_day']=pd.to_datetime(df['登录时点'],errors='raise').dt.normalize()
    if df.groupby('用户编号')['注册日期'].nunique().gt(1).any():raise ValueError('One user has conflicting registration dates')
    if df.login_day.lt(df['注册日期']).any():raise ValueError('Login predates registration')
    end=pd.Timestamp(observation_end).normalize()
    df=df.loc[df.login_day.le(end)].copy()
    if df.empty:raise ValueError('No observed users by the specified cutoff')
    users=df[['用户编号','注册日期']].drop_duplicates('用户编号')
    sizes=users.groupby('注册日期')['用户编号'].nunique().sort_index()
    days=df[['用户编号','注册日期','login_day']].drop_duplicates(['用户编号','login_day'])
    days['lag']=(days.login_day-days['注册日期']).dt.days
    observed=days.loc[days.lag.lt(horizon_days)].groupby(['注册日期','lag'])['用户编号'].nunique()
    counts=observed.unstack('lag',fill_value=0).reindex(index=sizes.index,columns=range(horizon_days),fill_value=0).astype(float)
    for lag in counts.columns:
        counts.loc[sizes.index+pd.Timedelta(days=lag)>end,lag]=np.nan
    rates=counts.div(sizes,axis=0)
    if (rates>1).any().any():raise AssertionError('Retention cannot exceed one after user-day deduplication')
    registered=users.groupby('注册日期')['用户编号'].nunique()
    active=days.groupby('login_day')['用户编号'].nunique()
    dates=pd.date_range(users['注册日期'].min(),end,freq='D')
    daily=pd.DataFrame({'registered_users':registered.reindex(dates,fill_value=0),'active_users':active.reindex(dates,fill_value=0)})
    daily.index.name='date'
    daily['registration_growth']=daily.registered_users.pct_change(fill_method=None).replace([np.inf,-np.inf],np.nan)
    overall=[]
    for lag in range(horizon_days):
        eligible=sizes.index+pd.Timedelta(days=lag)<=end
        denominator=int(sizes.loc[eligible].sum())
        numerator=int(counts[lag].sum())
        overall.append({'day':lag,'retained_users':numerator,'eligible_registered_users':denominator,
            'retention':numerator/denominator if denominator else np.nan,
            'textbook_fixed_denominator_retention':numerator/int(sizes.sum())})
    return daily,counts,rates,pd.DataFrame(overall),{'users':int(sizes.sum()),'user_day_rows':len(days),
        'cohorts':len(sizes),'observed_cohort_day_cells':int(counts.notna().sum().sum()),
        'not_yet_observed_cells':int(counts.isna().sum().sum())}


def main(argv=None):
    ctx=context('B07',argv)
    raw=pd.read_excel(ctx.source('cohort_raw'),sheet_name='用户登录明细',usecols='A:C')
    daily,counts,rates,overall,info=calculate(raw,OBSERVATION_END)
    wb=load_workbook(ctx.source('cohort_reference'),read_only=True,data_only=True)
    rows=list(wb['参考答案'].iter_rows(max_row=29,max_col=33,values_only=True))
    reference_daily=pd.DataFrame([{'date':r[0],'registered_users':r[1],'active_users':r[6]} for r in rows[3:12]]).set_index('date')
    checks={'campaign_daily':compare_numeric(daily.reindex(reference_daily.index)[['registered_users','active_users']],reference_daily[['registered_users','active_users']],'Campaign registration/active users',tolerance=0)}
    ref_cohorts=[r[0] for r in rows[19:28]]
    if list(rates.index)!=ref_cohorts:raise AssertionError('Cohort dates differ from the reference')
    expected=np.array([[np.nan if v is None else float(v) for v in r[1:32]] for r in rows[19:28]])
    # A blank pivot cell within an observed calendar window can mean zero activity;
    # a future cell remains missing, regardless of how the pivot rendered it.
    mask=counts.notna().to_numpy()
    expected_observed=np.nan_to_num(expected[mask],nan=0.)
    checks['observed_retention_cells']=compare_numeric(rates.to_numpy()[mask],expected_observed,'Observed cohort retention')
    checks['textbook_total_row']=compare_numeric(overall.textbook_fixed_denominator_retention,[float(v) for v in rows[28][1:32]],'Textbook fixed-denominator total row')
    wb.close()
    save_table(daily.reset_index(),ctx.output/'daily_activity.csv')
    save_table(counts.rename_axis(columns='day').reset_index().rename(columns={'注册日期':'registration_date'}),ctx.output/'retention_counts.csv')
    save_table(rates.rename_axis(columns='day').reset_index().rename(columns={'注册日期':'registration_date'}),ctx.output/'retention_rates.csv')
    save_table(overall,ctx.output/'overall_retention.csv')
    plt=plotting()
    from matplotlib.ticker import PercentFormatter
    fig,axes=plt.subplots(2,1,figsize=(14,9),height_ratios=[2,1],layout='constrained')
    shown=rates.loc[:,[0,1,2,3,7,14,21,22,23,24,25,26,27,28,29,30]]
    cmap=plt.get_cmap('YlGnBu').copy();cmap.set_bad('#e5e9eb')
    im=axes[0].imshow(np.ma.masked_invalid(shown.to_numpy()),aspect='auto',vmin=0,vmax=1,cmap=cmap)
    axes[0].set(xticks=range(len(shown.columns)),xticklabels=[f'D{x}' for x in shown.columns],
        yticks=range(len(shown)),yticklabels=[x.strftime('%m-%d') for x in shown.index],
        title='按注册日分组的日留存（关键日龄，灰色为未到期）',ylabel='注册日期')
    for (i,j),v in np.ndenumerate(shown.to_numpy()):
        if np.isfinite(v):axes[0].text(j,i,f'{v:.0%}',ha='center',va='center',fontsize=8,color='white' if v>.6 else '#17364d')
    fig.colorbar(im,ax=axes[0],format=PercentFormatter(xmax=1),fraction=.025,pad=.02)
    axes[1].plot(overall.day,overall.retention,color='#087f8c',label='只以达到观察日的群体为分母')
    axes[1].plot(overall.day,overall.textbook_fixed_denominator_retention,color='#9caeba',linestyle='--',label='教材总计：固定全部注册用户为分母')
    axes[1].yaxis.set_major_formatter(PercentFormatter(xmax=1))
    axes[1].set(title='汇总留存与观察成熟度',xlabel='注册后第几日',ylabel='留存率')
    axes[1].legend(fontsize=9)
    axes[1].text(.65,.40,f"D30：成熟群体 {overall.iloc[-1].retention:.3%}\n固定全部用户 {overall.iloc[-1].textbook_fixed_denominator_retention:.3%}",
                 transform=axes[1].transAxes,fontsize=10,color='#466070')
    fig.suptitle('同期群：同日重复登录去重，未到观察日不填零',fontsize=17,color='#17364d')
    fig.savefig(ctx.output/'cohort_retention.png',dpi=150)
    plt.close(fig)
    last=overall.iloc[-1]
    summary={'lab_id':'B07','status':'reference_verified','source_rows':len(raw),'observation_end':str(OBSERVATION_END.date()),
        'horizon_days':31,**info,'reference_checks':checks,'day_1_retention':float(overall.iloc[1].retention),
        'day_30_eligible_users':int(last.eligible_registered_users),'day_30_retention':float(last.retention),
        'day_30_textbook_fixed_denominator_retention':float(last.textbook_fixed_denominator_retention),
        'registration_scope':'Distinct users identifiable from the supplied login records; no independent registration ledger supplied.'}
    report=f"""# B07 同期群日留存

## 结果

资料有{len(raw):,}条登录记录、{info['users']:,}名可识别注册用户。同一用户同日登录去重后为{info['user_day_rows']:,}条用户日记录，分为{info['cohorts']}个注册日期群体。活动期为2020-07-01至07-09，观察截止2020-07-31。

![同期群留存](cohort_retention.png)

## 定义与参考核对

日留存=该注册群体在注册后第d日登录的不同用户数/该群体注册人数。首日为D0，次日为D1，不能把同日多次登录计为多人。注册日期不一致、登录早于注册会报错。

活动期间的每日注册/登录人数、{info['observed_cohort_day_cells']}个已观察的群体日留存单元，以及教材固定分母的31个总计值均通过对照。{info['not_yet_observed_cells']}个尚未到期的群体日格子保留为空；已到期但没有登录才填0。

## 对总计分母的补充

教材总计行将所有{info['users']}名注册用户固定作为每个日龄的分母。到D30，只有7月1日注册的{int(last.eligible_registered_users)}名用户达到观察日。固定分母结果为{last.textbook_fixed_denominator_retention:.4%}，按成熟群体计算为{last.retention:.4%}。两条口径都保存，后一条用于跨日龄汇总，避免把未到期群体算成流失。

## 解读范围

日留存不要求每天单调下降，用户可以间隔几天后回来。曲线变化本身不能证明促销或竞品活动的因果影响。当前分母来自登录资料可识别用户，没有独立注册总表，无法验证完全未登录用户是否被遗漏。完整31日群体矩阵和两种汇总口径见CSV聚合结果。
"""
    finish(ctx,summary,report)


if __name__=='__main__':main()
