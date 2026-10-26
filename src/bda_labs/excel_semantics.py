"""Small, explicit Excel-to-Python conventions; this is not a workbook formula engine."""
from decimal import Decimal,ROUND_HALF_UP
import calendar,math
import numpy as np
import pandas as pd
from openpyxl.utils.cell import coordinate_to_tuple,get_column_letter
from .tabular_tools import unique_key


def excel_round(value,digits=0):
    if isinstance(digits,bool) or int(digits)!=digits:raise ValueError('Integer decimal places required')
    number=Decimal(str(value))
    if not number.is_finite():raise ValueError('Finite number required')
    return float(number.quantize(Decimal(1).scaleb(-int(digits)),rounding=ROUND_HALF_UP))


def approximate_lookup(values,lower_bounds,labels):
    bounds=np.asarray(lower_bounds,dtype=float);x=np.asarray(values,dtype=float)
    if bounds.ndim!=1 or len(bounds)!=len(labels) or len(bounds)==0 or not np.isfinite(bounds).all() or np.any(np.diff(bounds)<=0):raise ValueError('Approximate lookup requires strictly increasing finite lower bounds')
    if not np.isfinite(x).all() or (x<bounds[0]).any():raise ValueError('No lookup band exists for an invalid/below-minimum value')
    return np.asarray(labels,dtype=object)[np.searchsorted(bounds,x,side='right')-1]


def exact_lookup(values,keys,labels):
    table=pd.DataFrame({'key':list(keys),'label':list(labels)})
    unique_key(table,'key')
    if table.label.isna().any():raise ValueError('A matched key needs a populated label')
    values=pd.Series(values)
    if values.isna().any():raise ValueError('Missing lookup keys require explicit handling')
    return values.map(table.set_index('key').label)


def offset_address(origin,rows,columns,height=1,width=1):
    if any(isinstance(v,bool) or not isinstance(v,(int,np.integer)) for v in [rows,columns,height,width]) or min(height,width)<1:raise ValueError('Integer offsets and positive dimensions required')
    r,c=coordinate_to_tuple(origin);r+=rows;c+=columns
    if r<1 or c<1 or r+height-1>1048576 or c+width-1>16384:raise ValueError('OFFSET extends outside the worksheet')
    return f'{get_column_letter(c)}{r}:{get_column_letter(c+width-1)}{r+height-1}'


def workday(start,n,holidays=()):
    if isinstance(n,bool) or not isinstance(n,(int,np.integer)):raise ValueError('Integer number of working days required')
    current=pd.Timestamp(start).normalize();excluded={pd.Timestamp(d).normalize() for d in holidays}
    direction=1 if n>=0 else -1;remaining=abs(int(n))
    while remaining:
        current+=pd.Timedelta(days=direction)
        if current.weekday()<5 and current not in excluded:remaining-=1
    return current


def networkdays(start,end,holidays=()):
    start,end=pd.Timestamp(start).normalize(),pd.Timestamp(end).normalize()
    if start>end:return -networkdays(end,start,holidays)
    excluded={pd.Timestamp(d).normalize() for d in holidays}
    return sum(d.weekday()<5 and d not in excluded for d in pd.date_range(start,end,freq='D'))


def tariff(quantity,thresholds,rates,policy):
    q=Decimal(str(quantity));bounds=[Decimal(str(x)) for x in thresholds];prices=[Decimal(str(x)) for x in rates]
    if not q.is_finite() or q<0 or len(prices)!=len(bounds)+1 or any(not v.is_finite() or v<0 for v in prices) or any(not v.is_finite() or v<=0 for v in bounds) or any(b<=a for a,b in zip(bounds,bounds[1:])):raise ValueError('Invalid quantity or tariff schedule')
    if policy=='whole_tier':amount=q*prices[sum(q>b for b in bounds)]
    elif policy=='progressive':
        previous=Decimal(0);amount=Decimal(0)
        for upper,price in zip([*bounds,q],prices):
            used=max(Decimal(0),min(q,upper)-previous);amount+=used*price;previous=upper
    else:raise ValueError('Choose whole_tier or progressive explicitly')
    return float(amount.quantize(Decimal('.01'),rounding=ROUND_HALF_UP))


def frequency(values,upper_bounds):
    bins=np.asarray(upper_bounds,dtype=float);values=np.asarray(values,dtype=float)
    if bins.ndim!=1 or not np.isfinite(bins).all() or not np.isfinite(values).all() or np.any(np.diff(bins)<=0):raise ValueError('Frequency bins must be finite and strictly increasing')
    return np.bincount(np.searchsorted(bins,values,side='left'),minlength=len(bins)+1)


def subject_passes(scores):
    values=np.asarray(scores,dtype=float)
    if values.ndim!=2 or values.shape[1]!=3 or not np.isfinite(values).all() or (values<0).any() or (values>100).any():raise ValueError('Expected three valid numeric subject scores')
    return (values>=60).all(axis=1),(values>=60).any(axis=1)
