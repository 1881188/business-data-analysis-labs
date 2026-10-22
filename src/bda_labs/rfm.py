"""B05: Excel-equivalent RFM at receipt/customer grain, with reference reconciliation."""
from __future__ import annotations

from datetime import datetime
from decimal import Decimal

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .common import context, finish, plotting, save_table

AS_OF = pd.Timestamp("2019-08-03")
SEGMENTS = {
    (True, True, True): "重要价值用户", (True, False, True): "重要发展用户",
    (False, True, True): "重要保持用户", (False, False, True): "重要挽留用户",
    (True, True, False): "一般价值用户", (True, False, False): "一般发展用户",
    (False, True, False): "一般保持用户", (False, False, False): "一般挽留用户",
}


def money_cents(values: pd.Series) -> pd.Series:
    """Reject sub-cent amounts instead of silently rounding unexplained source values."""
    def convert(value):
        amount = Decimal(str(value)) * 100
        if not amount.is_finite() or amount != amount.to_integral_value():
            raise ValueError(f"Amount is not an exact number of cents: {value!r}")
        return int(amount)
    return values.map(convert).astype("int64")


def calculate(raw: pd.DataFrame, as_of=AS_OF) -> pd.DataFrame:
    required = ["流水号", "用户编号", "销售日期", "销售金额"]
    if raw[required].isna().any().any():
        raise ValueError("RFM key/date/amount fields contain missing values; review the input.")
    df = raw[required].copy()
    df["用户编号"] = df["用户编号"].astype("string")
    df["销售日期"] = pd.to_datetime(df["销售日期"], errors="raise").dt.normalize()
    if (df["销售日期"] > as_of).any():
        raise ValueError("Sales after the analysis date must not enter RFM.")
    if df.groupby("流水号")["用户编号"].nunique().gt(1).any():
        raise ValueError("A receipt is associated with multiple customers.")
    df["amount_cents"] = money_cents(df["销售金额"])
    result = df.groupby("用户编号", sort=True).agg(
        last_purchase=("销售日期", "max"), frequency=("流水号", "nunique"),
        monetary_cents=("amount_cents", "sum"), line_count=("流水号", "size"))
    result["recency_days"] = (as_of - result["last_purchase"]).dt.days
    n = len(result)
    # Exact mean comparisons avoid boundary changes caused by floating point rounding.
    result["r_high"] = result["recency_days"] * n < result["recency_days"].sum()
    result["f_high"] = result["frequency"] * n > result["frequency"].sum()
    result["m_high"] = result["monetary_cents"] * n > result["monetary_cents"].sum()
    result["segment"] = [SEGMENTS[key] for key in zip(result.r_high, result.f_high, result.m_high)]
    return result.rename_axis("customer_id").reset_index()


def reference_table(path) -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    records = [r for r in wb["参考答案"].iter_rows(min_row=7, max_col=9, values_only=True)
               if r[0] is not None and isinstance(r[1], datetime)]
    wb.close()
    result = pd.DataFrame(records, columns=["customer_id", "last_purchase", "frequency", "monetary",
        "recency_days", "r_flag", "f_flag", "m_flag", "segment"])
    result["customer_id"] = result["customer_id"].astype("string")
    result["monetary_cents"] = money_cents(result["monetary"])
    return result


def reconcile(actual: pd.DataFrame, reference: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    cols = ["customer_id", "last_purchase", "frequency", "monetary_cents", "recency_days", "segment"]
    joined = actual[cols].merge(reference[cols], on="customer_id", how="outer", suffixes=("_python", "_excel"),
                                indicator=True, validate="one_to_one")
    metrics = {"customer_id_sets_match": bool(joined["_merge"].eq("both").all())}
    mismatch = ~joined["_merge"].eq("both")
    for col in cols[1:]:
        unequal = joined[f"{col}_python"].ne(joined[f"{col}_excel"])
        metrics[f"{col}_mismatches"] = int(unequal.sum())
        mismatch |= unequal
    return joined.loc[mismatch].drop(columns="_merge"), metrics


def main(argv=None):
    ctx = context("B05", argv)
    raw = pd.read_excel(ctx.source("rfm_raw"), sheet_name="原始数据", dtype={"用户编号": "string", "流水号": "string"})
    result = calculate(raw)
    reference = reference_table(ctx.source("rfm_reference"))
    differences, checks = reconcile(result, reference)
    save_table(result, ctx.output / "customers.csv")
    save_table(differences, ctx.output / "reference_differences.csv")
    if not differences.empty or not checks["customer_id_sets_match"]:
        raise AssertionError("RFM differs from the selected Excel reference; see reference_differences.csv")
    grouped = result.groupby("segment").agg(customers=("customer_id", "size"),
        receipts=("frequency", "sum"), monetary_cents=("monetary_cents", "sum")).sort_values("monetary_cents", ascending=False)
    grouped["customer_share"] = grouped.customers / len(result)
    grouped["revenue_share"] = grouped.monetary_cents / grouped.monetary_cents.sum()
    grouped["cumulative_revenue_share"] = grouped.revenue_share.cumsum()
    save_table(grouped.reset_index(), ctx.output / "segments.csv")
    cents = int(grouped.monetary_cents.sum())
    if cents != int(money_cents(raw["销售金额"]).sum()): raise AssertionError("Amount reconciliation failed")
    plt = plotting()
    fig, axes = plt.subplots(1, 2, figsize=(15, 6), layout="constrained")
    order = grouped.iloc[::-1]
    colors = ["#087f8c" if s == "重要价值用户" else "#6b8ba4" for s in order.index]
    axes[0].barh(order.index, order.customers, color=colors)
    axes[0].set(title="客户数量", xlabel="客户数（人）")
    axes[1].barh(order.index, order.revenue_share * 100, color=colors)
    axes[1].set(title="销售额贡献", xlabel="销售额占比（%）")
    for y, v in enumerate(order.revenue_share): axes[1].text(v*100 + .5, y, f"{v:.1%}", va="center", fontsize=10)
    axes[1].set_xlim(0, max(order.revenue_share) * 115)
    fig.suptitle("RFM客户分层：人数与金额贡献", fontsize=18, color="#17364d")
    fig.savefig(ctx.output / "rfm_segments.png", dpi=150)
    plt.close(fig)
    important = grouped.loc["重要价值用户"]
    summary = {"lab_id": "B05", "status": "reference_verified", "as_of": str(AS_OF.date()),
        "source_lines": len(raw), "customers": len(result), "receipts": int(result.frequency.sum()),
        "sales_cents": cents, "thresholds": {"recency_mean_days": float(result.recency_days.mean()),
        "frequency_mean": float(result.frequency.mean()), "monetary_mean": float(result.monetary_cents.mean()/100)},
        "reference_checks": checks, "segments": grouped.reset_index().to_dict("records")}
    report = f"""# B05 RFM客户价值分层

## 运行结果

原始数据{len(raw):,}条商品销售明细，合并为{int(result.frequency.sum()):,}笔不同流水号、{len(result):,}名客户，销售金额{cents/100:,.2f}元。分析时点固定为{AS_OF.date()}。

Python结果与Excel参考的{len(reference):,}名客户逐一比较：客户集合、最近消费日期、F、M、R、分层标签全部一致。金额使用整数分对账；原始销售表没有因重复外观而任意去重。

![RFM客户分层](rfm_segments.png)

重要价值用户{int(important.customers):,}人，占客户{important.customer_share:.2%}，贡献销售额{important.revenue_share:.2%}。可优先为这组用户设计服务与留存实验，但当前数据不能证明营销措施的增量效果。

## Excel到Python

- R = 分析日期减客户最后一次消费日期。R小于全体客户平均值才标为高。
- F = 每位客户不同流水号数量。原表一张小票可能多行商品，不能用行数或消费日期数代替。
- M = 全部销售明细金额之和。F、M严格大于客户均值标为高，相等归低。
- 八类标签对应教材业务背景工作表的分类表，等价于IF与VLOOKUP组合。

均值阈值：R={result.recency_days.mean():.8f}天，F={result.frequency.mean():.8f}笔，M={result.monetary_cents.mean()/100:.8f}元。

## 解释与限制

RFM是观察期内的规则分层，不是客户未来价值的因果预测。分层使用当前全体客户均值，阈值会随样本变化。参考核对使用教材已有缓存值，并由Python独立从原始明细计算，不声称已重算原Excel软件。

本地完整输出包括customers.csv、segments.csv、reference_differences.csv和summary.json。公开报告只展示聚合结果，不公开客户级明细。
"""
    finish(ctx, summary, report)


if __name__ == "__main__": main()
